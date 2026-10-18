"""文件解析：docx / md / txt 提取文本，供对话分析。"""
from pathlib import Path

# 提取文本的最大字符数（避免超长文件撑爆上下文）
MAX_CHARS = 50000


def parse_file(path: str | Path) -> dict:
    """解析文件，返回 {filename, text, char_count, truncated}。"""
    p = Path(path)
    filename = p.name
    ext = p.suffix.lower()

    if ext == ".docx":
        text = _parse_docx(p)
    elif ext in (".md", ".txt", ".log", ".text", ".markdown"):
        text = _parse_text(p)
    elif ext == ".doc":
        # 旧版 .doc 不支持（二进制），提示转 docx
        raise ValueError("暂不支持 .doc 格式，请另存为 .docx 后再上传")
    else:
        raise ValueError(f"不支持的文件格式：{ext}，请上传 docx / md / txt")

    text = text.strip()
    truncated = len(text) > MAX_CHARS
    if truncated:
        text = text[:MAX_CHARS]

    return {
        "filename": filename,
        "text": text,
        "char_count": len(text),
        "truncated": truncated,
    }


def _parse_docx(p: Path) -> str:
    from docx import Document
    doc = Document(str(p))
    parts = []
    for para in doc.paragraphs:
        t = para.text.strip()
        if t:
            parts.append(t)
    # 表格
    for tbl in doc.tables:
        for row in tbl.rows:
            cells = [c.text.strip() for c in row.cells]
            line = " | ".join(cells).strip()
            if line and line != "|":
                parts.append(line)
    return "\n".join(parts)


def _parse_text(p: Path) -> str:
    # 尝试常见编码
    raw = p.read_bytes()
    for enc in ("utf-8", "gbk", "gb18030", "utf-16"):
        try:
            return raw.decode(enc)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return raw.decode("utf-8", errors="ignore")
